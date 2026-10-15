"""Resume document I/O (``.docx`` only, via python-docx).

``extract_text`` reads an uploaded resume into plain text for the agents.
Legacy binary ``.doc`` files are not supported — convert them to ``.docx`` first.
"""

import math
from io import BytesIO
from typing import List, Optional, Tuple

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph


def extract_text(data: bytes) -> str:
    """Extract plain text from ``.docx`` bytes (paragraphs and table cells)."""
    document = Document(BytesIO(data))

    lines: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    # Include table content (many resumes use tables for layout).
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines)


def resume_docx_path():
    """Original upload, stored beside the database."""
    from ..config import DB_PATH

    return DB_PATH.parent / "resume.docx"


def save_resume_docx(data: bytes) -> None:
    path = resume_docx_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


def load_resume_docx() -> Optional[bytes]:
    path = resume_docx_path()
    if path.is_file():
        return path.read_bytes()
    return None


def paragraph_count(data: bytes) -> int:
    return len(Document(BytesIO(data)).paragraphs)


def _iter_paragraphs(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_in_run(paragraph: Paragraph, original: str, replacement: str) -> bool:
    """Replace only when the span sits inside one run, so bold survives."""
    for run in paragraph.runs:
        if original in (run.text or ""):
            run.text = run.text.replace(original, replacement, 1)
            return True
    return False


def _length_emu(value, default: int = 0) -> int:
    if value is None:
        return default
    return int(value)


def _font_emu(run, paragraph) -> int:
    size = run.font.size
    if size is None and paragraph.style is not None:
        size = paragraph.style.font.size
    if size is None:
        size = Pt(11)
    return int(size)


def _paragraph_spacing(paragraph) -> Tuple[int, int, Optional[float], Optional[int]]:
    fmt = paragraph.paragraph_format
    before = fmt.space_before
    after = fmt.space_after
    spacing = fmt.line_spacing
    if paragraph.style is not None:
        style_fmt = paragraph.style.paragraph_format
        if before is None:
            before = style_fmt.space_before
        if after is None:
            after = style_fmt.space_after
        if spacing is None:
            spacing = style_fmt.line_spacing
    exact = None
    multiple = None
    if isinstance(spacing, (int, float)) and not isinstance(spacing, bool):
        if isinstance(spacing, float) or (isinstance(spacing, int) and spacing < 20):
            multiple = float(spacing)
        else:
            exact = int(spacing)
    elif spacing is not None:
        exact = int(spacing)
    return _length_emu(before), _length_emu(after), multiple, exact


def _paragraph_height(paragraph: Paragraph, content_width: int, char_factor: float = 0.5) -> int:
    """Wrapped height of one paragraph.

    ``char_factor`` is the assumed glyph width in ems. 0.5 is deliberately
    wide so a normal estimate does not miss a page break.
    """
    before, after, multiple, exact = _paragraph_spacing(paragraph)
    left = _length_emu(paragraph.paragraph_format.left_indent)
    right = _length_emu(paragraph.paragraph_format.right_indent)
    width = max(content_width - left - right, content_width // 4)
    fonts: List[int] = []
    text_width = 0
    for run in paragraph.runs:
        font = _font_emu(run, paragraph)
        fonts.append(font)
        # Average Latin glyph is a bit under half an em. 0.5 keeps the estimate
        # on the safe side of a real page break.
        text_width += len(run.text or "") * font * char_factor
    font = max(fonts) if fonts else int(Pt(11))
    if text_width == 0 and paragraph.text:
        text_width = len(paragraph.text) * font * char_factor
    if multiple is not None:
        line = int(font * multiple)
    elif exact is not None:
        line = exact
    else:
        line = int(font * 1.0)
    line = max(line, int(font))
    if text_width <= 0:
        lines = 1 if (paragraph.text or before or after) else 0
    else:
        lines = max(1, math.ceil(text_width / width))
    return before + lines * line + after


def _body_height(document: Document, content_width: int, char_factor: float = 0.5) -> int:
    height = 0
    for paragraph in document.paragraphs:
        height += _paragraph_height(paragraph, content_width, char_factor)
    for table in document.tables:
        for row in table.rows:
            seen = set()
            row_height = 0
            cells = []
            for cell in row.cells:
                marker = id(cell._tc)
                if marker in seen:
                    continue
                seen.add(marker)
                cells.append(cell)
            share = max(1, len(cells))
            for cell in cells:
                cell_width = _length_emu(cell.width, content_width // share)
                cell_height = sum(
                    _paragraph_height(paragraph, cell_width, char_factor)
                    for paragraph in cell.paragraphs
                )
                row_height = max(row_height, cell_height)
            height += row_height
    return height


def _section_box(data: bytes) -> Tuple[Document, int, int]:
    document = Document(BytesIO(data))
    section = document.sections[0]
    content_width = int(section.page_width - section.left_margin - section.right_margin)
    content_height = int(section.page_height - section.top_margin - section.bottom_margin)
    return document, content_width, content_height


def page_metrics(data: bytes, char_factor: float = 0.5) -> Tuple[int, int]:
    """Return ``(page_count, body_height)`` from section size and wrapped text.

    This is a layout estimate, not Word's own pagination. The default glyph
    width is biased high so a normal call does not miss a page break.
    """
    document, content_width, content_height = _section_box(data)
    if content_width <= 0 or content_height <= 0:
        return 1, 0
    body = _body_height(document, content_width, char_factor)
    pages = max(1, math.ceil(body / content_height))
    return pages, body


def one_page_factor(data: bytes) -> float:
    """Largest glyph width that still lays the upload out as one page.

    A packed one-page résumé estimates as two pages at the default width.
    Rewrites are judged with the widest factor that still fits the upload,
    so an edit that adds a line no longer hides inside that overestimate.
    """
    high = 0.5
    if page_metrics(data, high)[0] <= 1:
        return high
    low = 0.15
    best = low
    for _ in range(14):
        mid = (low + high) / 2
        if page_metrics(data, mid)[0] <= 1:
            best = mid
            low = mid
        else:
            high = mid
    return best


def _one_page_budget(original: bytes) -> Tuple[float, int]:
    """Return ``(glyph factor, maximum body height)`` for a one-page result."""
    factor = one_page_factor(original)
    _pages, body = page_metrics(original, factor)
    _document, _width, content_height = _section_box(original)
    if content_height > 0 and body <= content_height:
        return factor, content_height
    return factor, body


def estimated_pages(data: bytes) -> int:
    return page_metrics(data)[0]


def keep_one_page(
    original: bytes, current: bytes, patches: List[dict]
) -> Tuple[List[dict], List[dict]]:
    """Drop patches that would make the finished résumé longer than one page.

    The uploaded file is the one-page reference. A later try may use space a
    shorter edit freed, but the result may not grow past one page. A patch
    that does not land in the file is kept only when its replacement is no
    longer than the original span.
    """
    factor, budget = _one_page_budget(original)
    working = current
    kept: List[dict] = []
    dropped: List[dict] = []
    for patch in patches:
        trial, applied = apply_patches_to_docx(working, [[patch]])
        changed = bool(applied and applied[0].get("applied"))
        if not changed:
            original_text = str(patch.get("original") or "")
            replacement = str(patch.get("replacement") or "")
            if len(replacement) > len(original_text) or "\n" in replacement:
                dropped.append({**patch, "reason": "overflows_page"})
            else:
                kept.append(patch)
            continue
        if page_metrics(trial, factor)[1] <= budget:
            kept.append(patch)
            working = trial
        else:
            dropped.append({**patch, "reason": "overflows_page"})
    return kept, dropped


def enforce_one_page(
    original: bytes, patch_lists: List[List[dict]]
) -> Tuple[bytes, List[dict]]:
    """Apply a patch chain and skip any edit that would spill onto page two."""
    factor, budget = _one_page_budget(original)
    working = original
    results: List[dict] = []
    for patches in patch_lists:
        for patch in patches:
            trial, applied = apply_patches_to_docx(working, [[patch]])
            row = dict(patch)
            changed = bool(applied and applied[0].get("applied"))
            if changed and page_metrics(trial, factor)[1] <= budget:
                working = trial
                row["applied"] = True
            else:
                row["applied"] = False
                if changed:
                    row["reason"] = "overflows_page"
            results.append(row)
    return working, results


def apply_patches_to_docx(
    data: bytes, patch_lists: List[List[dict]]
) -> Tuple[bytes, List[dict]]:
    """Apply patch lists in order. A span split across runs is left unapplied.

    ``paragraph.text`` is never assigned: that property rebuilds the paragraph
    and drops bold and other run formatting.
    """
    document = Document(BytesIO(data))
    results: List[dict] = []
    for patches in patch_lists:
        for patch in patches:
            original = str(patch.get("original") or "")
            replacement = str(patch.get("replacement") or "")
            applied = False
            if original and original != replacement:
                for paragraph in _iter_paragraphs(document):
                    if original in paragraph.text and _replace_in_run(
                        paragraph, original, replacement
                    ):
                        applied = True
                        break
            row = dict(patch)
            row["applied"] = applied
            results.append(row)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue(), results
